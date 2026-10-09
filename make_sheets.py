#!/usr/bin/env python3
"""Generate the Northwind Stays UAT test sheets.

    python make_sheets.py --base-url https://USERNAME.github.io/northwind-stays

Writes:
  sheets/Northwind_Tests_Build_1.0.0.xlsx       rows open {base}/v1/...
  sheets/Northwind_Tests_Build_1.1.0.xlsx       the same rows, opening {base}/v2/...
  sheets/Northwind_NewFeature_Rows_1.1.0.xlsx   Free cancellation rows on {base}/v2/...
  verify/test_rows.json                          the same rows as structured steps for verify/verify.js

Every expected value (counts, names, totals) is computed here from src/data.js
using Build 1.0.0 behaviour. Nothing is typed in by hand.
"""

import argparse
import datetime
import json
import sys
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent

COLUMNS = [
    "StepID", "Persona", "Feature", "WhatToDo", "TestData", "ExpectedResult",
    "Impact", "Likelihood", "RiskScore", "Smoke", "Status", "RunDate", "Notes",
]
WIDTHS = {
    "StepID": 9, "Persona": 18, "Feature": 20, "WhatToDo": 70, "TestData": 34,
    "ExpectedResult": 60, "Impact": 9, "Likelihood": 11, "RiskScore": 10,
    "Smoke": 8, "Status": 10, "RunDate": 13, "Notes": 34,
}
VALIDATION_LAST_ROW = 500

TESTER = "tester@northwind.test"
LOCKED = "locked@northwind.test"
PASSWORD = "Test@1234"
OTP = "123456"
GUEST_NAME = "Alex Tester"
GUEST_EMAIL = "alex.tester@northwind.test"
GUEST_MOBILE = "91234567"

LOGIN_HEADING = "Log in to Northwind Stays"
BOOK_HEADING = "Complete your booking"


# --------------------------------------------------------------------------
# Data and Build 1.0.0 behaviour
# --------------------------------------------------------------------------

def load_data():
    text = (ROOT / "src" / "data.js").read_text(encoding="utf-8")
    return json.loads(text[text.index("{"): text.rindex("}") + 1])


DATA = load_data()
PROPS = DATA["properties"]
BY_NAME = {p["name"]: p for p in PROPS}


def prop(name):
    return BY_NAME[name]


def check_data_rules():
    """Fail fast if data.js no longer meets the rules the tests depend on."""
    problems = []
    cities = DATA["cities"]
    if len(PROPS) != 16:
        problems.append("expected 16 properties, found %d" % len(PROPS))
    if sorted({p["city"] for p in PROPS}) != sorted(cities) or len(cities) != 6:
        problems.append("properties must cover exactly the 6 cities")
    singapore = [p for p in PROPS if p["city"] == "Singapore"]
    if len(singapore) < 3 or len({p["price"] for p in singapore}) != len(singapore):
        problems.append("need at least 3 Singapore properties with different prices")
    if not any(p["price"] == 200 for p in PROPS):
        problems.append("need a property priced exactly S$200")
    for city in cities:
        prices = [p["price"] for p in PROPS if p["city"] == city]
        if len(prices) != len(set(prices)):
            problems.append("two properties in %s share a price" % city)
    for p in PROPS:
        if not (90 <= p["price"] <= 480 and 3 <= p["stars"] <= 5 and 7.0 <= p["rating"] <= 9.6 and 2 <= p["maxGuests"] <= 6):
            problems.append("%s is outside the allowed ranges" % p["name"])
        if p["type"] not in ("Hotel", "Serviced apartment") or not isinstance(p["freeCancellation"], bool):
            problems.append("%s has a bad type or freeCancellation value" % p["name"])
        if any(a not in DATA["amenities"] for a in p["amenities"]):
            problems.append("%s has an unknown amenity" % p["name"])
    if problems:
        sys.exit("data.js problems:\n  " + "\n  ".join(problems))


def parse_date(text):
    day, month, year = (int(x) for x in text.split("/"))
    return datetime.date(year, month, day)


class Stay:
    def __init__(self, city, checkin, checkout, guests=2, rooms=1):
        self.city, self.checkin, self.checkout = city, checkin, checkout
        self.guests, self.rooms = guests, rooms
        self.nights = (parse_date(checkout) - parse_date(checkin)).days
        assert self.nights > 0

    def query(self, **first):
        values = list(first.items()) + [
            ("city", self.city), ("checkin", self.checkin), ("checkout", self.checkout),
            ("guests", self.guests), ("rooms", self.rooms),
        ]
        return "&".join("%s=%s" % (k, quote(str(v), safe="/")) for k, v in values if v not in ("", None))

    def describe(self):
        return "%s, %s to %s, %d guest%s, %d room%s" % (
            self.city or "All cities", self.checkin, self.checkout,
            self.guests, "" if self.guests == 1 else "s", self.rooms, "" if self.rooms == 1 else "s")


def search(stay):
    return [p for p in PROPS if (not stay.city or p["city"] == stay.city) and p["maxGuests"] >= stay.guests]


def apply_filters(props, min_price=None, max_price=None, stars=(), types=(), amenities=(), free_cancellation=False):
    out = []
    for p in props:
        if min_price is not None and p["price"] < min_price:
            continue
        if max_price is not None and p["price"] > max_price:
            continue
        if stars and p["stars"] not in stars:
            continue
        if types and p["type"] not in types:
            continue
        if any(a not in p["amenities"] for a in amenities):
            continue
        if free_cancellation and not p["freeCancellation"]:
            continue
        out.append(p)
    return out


SORT_KEYS = {
    "Recommended": None,
    "Price (low to high)": lambda p: p["price"],
    "Price (high to low)": lambda p: -p["price"],
    "Rating (high to low)": lambda p: -p["rating"],
}


def sort_props(props, mode):
    key = SORT_KEYS[mode]
    ordered = sorted(props, key=lambda p: PROPS.index(p))  # data order
    return ordered if key is None else sorted(ordered, key=key)  # sorted() is stable: ties keep data order


def count_text(n):
    if n == 0:
        return "No properties match your filters."
    return "1 property found" if n == 1 else "%d properties found" % n


def money(amount):
    return "S${:,}".format(int(amount))


def stay_total(p, stay):
    return p["price"] * stay.nights * stay.rooms


def names(props):
    return [p["name"] for p in props]


# --------------------------------------------------------------------------
# Steps and checks (structured, rendered to English and to JSON)
# --------------------------------------------------------------------------

def fill(label, value):
    return {"op": "fill", "label": label, "value": value}


def select(label, option):
    return {"op": "select", "label": label, "option": option}


def tick(label):
    return {"op": "tick", "label": label}


def click(name, role="button", kind=None, nav=False):
    return {"op": "click", "name": name, "role": role, "kind": kind or role, "nav": nav}


def click_card(card, name="View details"):
    return {"op": "click_card", "card": card, "name": name, "role": "link", "kind": "button", "nav": True}


def checkpoint(say, lead, *checks):
    return {"op": "check", "say": say, "lead": lead, "checks": list(checks)}


def heading(text):
    return {"type": "heading", "text": text}


def shows(text):
    return {"type": "shows", "text": text}


def not_shows(text):
    return {"type": "not_shows", "text": text}


def count_is(n):
    return {"type": "count", "text": count_text(n)}


def results_are(props):
    return {"type": "results", "names": names(props)}


def first_last(props):
    return {"type": "first_last", "first": props[0]["name"], "last": props[-1]["name"]}


def unticked(label):
    return {"type": "checked", "label": label, "value": False}


def button(name, present=True):
    return {"type": "button", "name": name, "present": present}


def card_badge(card, present=True):
    return {"type": "card_badge", "card": card, "text": "Free cancellation", "present": present}


def on_page(path):
    return {"type": "page", "path": path}


def field_visible(label):
    return {"type": "field", "label": label}


def join_names(items):
    items = list(items)
    if len(items) == 1:
        return items[0]
    return ", ".join(items[:-1]) + " and " + items[-1]


def clause(step):
    op = step["op"]
    if op == "fill":
        if step["value"] == "":
            return "leave the %s field empty" % step["label"]
        return "type %s into the %s field" % (step["value"], step["label"])
    if op == "select":
        return "select %s in the %s dropdown" % (step["option"], step["label"])
    if op == "tick":
        return "tick the %s checkbox" % step["label"]
    if op == "click":
        return "click the %s %s" % (step["name"], step["kind"])
    if op == "click_card":
        return "click the %s button on the %s card" % (step["name"], step["card"])
    if op == "check":
        return step["say"]
    raise ValueError(op)


def sentence(steps, lead=None):
    parts = [clause(s) for s in steps]
    if len(parts) == 1:
        body = parts[0]
    else:
        body = ", ".join(parts[:-1]) + ", then " + parts[-1]
    text = "%s, %s" % (lead, body) if lead else body
    return text[0].upper() + text[1:] + "."


def phrase(check):
    t = check["type"]
    if t == "heading":
        return 'the page heading is "%s"' % check["text"]
    if t == "shows":
        return '"%s" is shown on the page' % check["text"]
    if t == "not_shows":
        return '"%s" is not shown anywhere on the page' % check["text"]
    if t == "count":
        if check["text"] == count_text(0):
            return 'the page shows "%s"' % check["text"]
        return 'the results line reads "%s"' % check["text"]
    if t == "results":
        if not check["names"]:
            return "no property cards are shown"
        if len(check["names"]) == 1:
            return "the only property card shown is %s" % check["names"][0]
        return "exactly %d property cards are shown (in any order): %s" % (len(check["names"]), join_names(check["names"]))
    if t == "first_last":
        return "the first property card is %s and the last property card is %s" % (check["first"], check["last"])
    if t == "checked":
        return "the %s checkbox is %s" % (check["label"], "ticked" if check["value"] else "unticked")
    if t == "button":
        return ('a "%s" button is shown' if check["present"] else 'no "%s" button is shown') % check["name"]
    if t == "card_badge":
        if check["present"]:
            return 'the %s card shows a "%s" badge' % (check["card"], check["text"])
        return 'the %s card does not show a "%s" badge' % (check["card"], check["text"])
    if t == "page":
        return "the address bar shows %s" % check["path"]
    if t == "field":
        return "the %s field is visible" % check["label"]
    raise ValueError(t)


def group_text(lead, checks):
    out = []
    for i, check in enumerate(checks):
        text = phrase(check)
        if i == 0 and lead:
            text = "%s, %s" % (lead, text)
        out.append(text[0].upper() + text[1:] + ("" if text.endswith('."') else "."))
    return " ".join(out)


class Row:
    def __init__(self, persona, feature, start, sentences, expect, impact, likelihood,
                 smoke=False, stay=None, open_note="", final_lead=None, fails_on=()):
        self.persona, self.feature, self.start = persona, feature, start
        self.sentences = sentences  # list of (lead, [steps])
        self.expect = expect        # checks verified at the end
        self.final_lead = final_lead
        self.impact, self.likelihood, self.smoke = impact, likelihood, smoke
        self.stay, self.open_note = stay, open_note
        self.fails_on = list(fails_on)  # used only by the verifier, never written to a sheet
        self.step_id = None

    def steps(self):
        return [step for _, steps in self.sentences for step in steps]

    def what_to_do(self, base, build):
        url = "%s/%s/%s" % (base, build, self.start)
        text = "Open %s%s." % (url, self.open_note)
        for lead, steps in self.sentences:
            text += " " + sentence(steps, lead)
        return text

    def test_data(self):
        parts = []
        if self.stay:
            parts.append("Search: " + self.stay.describe())
        for step in self.steps():
            if step["op"] == "fill":
                parts.append("%s: %s" % (step["label"], step["value"] or "(empty)"))
            elif step["op"] == "select":
                parts.append("%s: %s" % (step["label"], step["option"]))
            elif step["op"] == "tick":
                parts.append("Tick: %s" % step["label"])
        return "; ".join(parts) if parts else "None"

    def expected_result(self):
        groups = [(s["lead"], s["checks"]) for s in self.steps() if s["op"] == "check"]
        groups.append((self.final_lead, self.expect))
        return " ".join(group_text(lead, checks) for lead, checks in groups)

    def to_json(self, base, build):
        steps = [{"op": "open", "path": self.start}] + self.steps()
        return {
            "id": self.step_id, "feature": self.feature, "smoke": self.smoke,
            "whatToDo": self.what_to_do(base, build), "expectedResult": self.expected_result(),
            "steps": steps, "checks": self.expect, "failsOn": self.fails_on,
        }


def as_tester(path_and_query):
    page, _, query = path_and_query.partition("?")
    return "%s?as=tester%s" % (page, "&" + query if query else "")


def results_url(stay):
    return as_tester("results.html?" + stay.query())


def details_url(p, stay):
    return as_tester("details.html?" + stay.query(id=p["id"]))


def book_url(p, stay):
    return as_tester("book.html?" + stay.query(id=p["id"]))


def guest_details(name=GUEST_NAME, email=GUEST_EMAIL, mobile=GUEST_MOBILE):
    return [fill("Full name", name), fill("Email", email), fill("Mobile number (+65)", mobile)]


def summary_shows(*texts):
    return [shows(t) for t in texts]


# --------------------------------------------------------------------------
# Rows
# --------------------------------------------------------------------------

def core_rows():
    rows = []
    sg = Stay("Singapore", "12/03/2027", "15/03/2027")
    everywhere = Stay("", "12/03/2027", "15/03/2027")
    sg_results = search(sg)
    login_steps = [fill("Email", TESTER), fill("Password", PASSWORD), click("Log in", nav=True)]

    # Login, one-time code, access control, logout
    rows.append(Row(
        "Registered guest", "Login and one-time code", "login.html",
        [(None, login_steps),
         ("on the one-time code page", [fill("One-time code", OTP), click("Verify", nav=True)])],
        [heading("Find a stay"), shows("Signed in as " + TESTER)],
        3, 3, smoke=True, final_lead="After clicking Verify"))
    rows.append(Row(
        "Registered guest", "Login", "login.html",
        [(None, [fill("Email", TESTER), fill("Password", "Wrong@1234"), click("Log in")])],
        [shows("Incorrect email or password."), heading(LOGIN_HEADING)],
        3, 2))
    rows.append(Row(
        "Locked user", "Login", "login.html",
        [(None, [fill("Email", LOCKED), fill("Password", PASSWORD), click("Log in")])],
        [shows("This account is locked. Please contact support."), heading(LOGIN_HEADING)],
        2, 2))
    rows.append(Row(
        "Registered guest", "One-time code", "login.html",
        [(None, login_steps),
         ("on the one-time code page", [fill("One-time code", "654321"), click("Verify")])],
        [shows("That code is incorrect. Please try again."), heading("Enter your one-time code")],
        3, 2))
    rows.append(Row(
        "Signed-out visitor", "Access control", "search.html",
        [], [heading(LOGIN_HEADING), on_page("login.html"), field_visible("Email"), field_visible("Password")],
        3, 1, open_note=" without logging in first and wait for the page to finish loading",
        final_lead="Instead of the search page"))
    rows.append(Row(
        "Registered guest", "Logout", "search.html?as=tester",
        [(None, [click("Log out", nav=True)])],
        [shows("You have been logged out."), heading(LOGIN_HEADING), on_page("login.html")],
        3, 2, smoke=True, final_lead="After clicking Log out"))

    # Search validation
    def search_steps(city, checkin, checkout, guests=None, rooms=None, nav=False):
        steps = [select("City", city), fill("Check-in (DD/MM/YYYY)", checkin), fill("Check-out (DD/MM/YYYY)", checkout)]
        if guests is not None:
            steps.append(select("Guests", str(guests)))
        if rooms is not None:
            steps.append(select("Rooms", str(rooms)))
        steps.append(click("Search", nav=nav))
        return steps

    rows.append(Row(
        "Registered guest", "Search validation", "search.html?as=tester",
        [(None, search_steps("Singapore", "2027-03-12", "15/03/2027"))],
        [shows("Enter dates as DD/MM/YYYY."), heading("Find a stay")],
        2, 3))
    rows.append(Row(
        "Registered guest", "Search validation", "search.html?as=tester",
        [(None, search_steps("Singapore", "15/03/2027", "12/03/2027"))],
        [shows("Check-out must be after check-in."), heading("Find a stay")],
        2, 2))
    rows.append(Row(
        "Registered guest", "Search validation", "search.html?as=tester",
        [(None, search_steps("Singapore", "01/03/2025", "03/03/2025"))],
        [shows("Check-in date cannot be in the past."), heading("Find a stay")],
        2, 2))

    # Search
    rows.append(Row(
        "Registered guest", "Search", "search.html?as=tester",
        [(None, search_steps("Singapore", sg.checkin, sg.checkout, 2, 1, nav=True))],
        [heading("Search results"), count_is(len(sg_results)), results_are(sg_results)],
        3, 3, smoke=True, final_lead="After clicking Search"))
    big = Stay("Singapore", "12/03/2027", "15/03/2027", guests=5)
    big_results = search(big)
    assert len(big_results) < len(sg_results)
    rows.append(Row(
        "Registered guest", "Search", "search.html?as=tester",
        [(None, search_steps("Singapore", big.checkin, big.checkout, big.guests, nav=True))],
        [count_is(len(big_results)), results_are(big_results)],
        3, 2, final_lead="After clicking Search"))

    # Filters, one at a time
    apply_btn = click("Apply filters")
    minimum = apply_filters(sg_results, min_price=200)
    assert any(p["price"] == 200 for p in minimum)
    rows.append(Row(
        "Registered guest", "Filter: min price", results_url(sg),
        [(None, [fill("Min price per night (S$)", "200"), apply_btn])],
        [count_is(len(minimum)), results_are(minimum)],
        2, 2, stay=sg, final_lead="After clicking Apply filters"))

    boundary = apply_filters(sg_results, max_price=200)
    assert any(p["price"] == 200 for p in boundary), "the max price boundary row needs a property at exactly S$200"
    rows.append(Row(
        "Registered guest", "Filter: max price", results_url(sg),
        [(None, [fill("Max price per night (S$)", "200"), apply_btn])],
        [count_is(len(boundary)), results_are(boundary)],
        2, 3, stay=sg, final_lead="After clicking Apply filters", fails_on=["v2"]))

    bkk = Stay("Bangkok", "12/03/2027", "15/03/2027")
    bkk_results = search(bkk)
    under = apply_filters(bkk_results, max_price=200)
    assert not any(p["price"] == 200 for p in bkk_results) and 0 < len(under) < len(bkk_results)
    rows.append(Row(
        "Registered guest", "Filter: max price", results_url(bkk),
        [(None, [fill("Max price per night (S$)", "200"), apply_btn])],
        [count_is(len(under)), results_are(under)],
        2, 2, stay=bkk, final_lead="After clicking Apply filters"))

    all_results = search(everywhere)
    five = apply_filters(all_results, stars=[5])
    rows.append(Row(
        "Registered guest", "Filter: star rating", results_url(everywhere),
        [(None, [tick("5 stars"), apply_btn])],
        [count_is(len(five)), results_are(five)],
        2, 2, stay=everywhere, final_lead="After clicking Apply filters"))

    serviced = apply_filters(sg_results, types=["Serviced apartment"])
    rows.append(Row(
        "Registered guest", "Filter: property type", results_url(sg),
        [(None, [tick("Serviced apartment"), apply_btn])],
        [count_is(len(serviced)), results_are(serviced)],
        2, 1, stay=sg, final_lead="After clicking Apply filters"))

    pets = apply_filters(all_results, amenities=["Pet friendly"])
    rows.append(Row(
        "Registered guest", "Filter: amenities", results_url(everywhere),
        [(None, [tick("Pet friendly"), apply_btn])],
        [count_is(len(pets)), results_are(pets)],
        2, 2, stay=everywhere, final_lead="After clicking Apply filters"))

    pool_gym = apply_filters(all_results, amenities=["Pool", "Gym"])
    assert len(pool_gym) < len(apply_filters(all_results, amenities=["Pool"]))
    rows.append(Row(
        "Registered guest", "Filter: amenities", results_url(everywhere),
        [(None, [tick("Pool"), tick("Gym"), apply_btn])],
        [count_is(len(pool_gym)), results_are(pool_gym)],
        2, 2, stay=everywhere, final_lead="After clicking Apply filters (only properties with both Pool and Gym)"))

    none_left = apply_filters(sg_results, stars=[3], amenities=["Pool"])
    assert not none_left
    rows.append(Row(
        "Registered guest", "Filter: combined", results_url(sg),
        [(None, [tick("3 stars"), tick("Pool"), apply_btn])],
        [count_is(0), results_are([])],
        1, 2, stay=sg, final_lead="After clicking Apply filters"))

    narrowed = apply_filters(sg_results, stars=[4], amenities=["Pool"])
    pool_only = apply_filters(sg_results, amenities=["Pool"])
    assert len(narrowed) != len(sg_results) and len(pool_only) != len(sg_results)
    rows.append(Row(
        "Registered guest", "Filter: clear filters", results_url(sg),
        [(None, [tick("4 stars"), tick("Pool"), apply_btn,
                 checkpoint("check the results line", "After clicking Apply filters", count_is(len(narrowed)))]),
         ("next", [click("Clear filters")])],
        [unticked("4 stars"), unticked("Pool"), count_is(len(sg_results)), results_are(sg_results)],
        2, 2, stay=sg, final_lead="After clicking Clear filters", fails_on=["v2"]))

    # Sorting
    rows.append(Row(
        "Registered guest", "Sort", results_url(sg),
        [(None, [select("Sort by", "Price (high to low)"), select("Sort by", "Recommended")])],
        [first_last(sort_props(sg_results, "Recommended"))],
        1, 1, stay=sg, final_lead="With Recommended selected"))
    for mode, impact, likelihood in (("Price (low to high)", 2, 2), ("Price (high to low)", 2, 1), ("Rating (high to low)", 1, 2)):
        rows.append(Row(
            "Registered guest", "Sort", results_url(sg),
            [(None, [select("Sort by", mode)])],
            [first_last(sort_props(sg_results, mode))],
            impact, likelihood, stay=sg, final_lead="After choosing " + mode))

    # Property details
    lantern = prop("Lantern Quay Hotel")
    rows.append(Row(
        "Registered guest", "Property details", results_url(sg),
        [(None, [click_card(lantern["name"])])],
        [heading(lantern["name"]), shows(money(lantern["price"]) + " per night"), shows("%d stars" % lantern["stars"]),
         shows("%.1f/10" % lantern["rating"]), shows("Nights: %d" % sg.nights), button("Book now")],
        3, 2, smoke=True, stay=sg, final_lead="On the details page"))

    # Booking
    confirm = click("Confirm booking", nav=True)
    marina = prop("Northwind Marina Suites")
    marina_total = money(stay_total(marina, sg))
    rows.append(Row(
        "Registered guest", "Booking", details_url(marina, sg),
        [(None, [click("Book now", role="link", kind="button", nav=True)]),
         ("on the booking page", [checkpoint(
             "check the booking summary", "Before confirming, on the booking page",
             *summary_shows("Nights: %d" % sg.nights, "Rooms: 1", "Total: " + marina_total, "Payment: Pay at property"))]
          + guest_details() + [confirm])],
        [heading("Booking confirmed"), shows(marina["name"]), shows("Total: " + marina_total), shows("Payment: Pay at property")],
        3, 3, smoke=True, stay=sg, final_lead="After clicking Confirm booking"))

    two_rooms = Stay("Singapore", "12/03/2027", "15/03/2027", guests=2, rooms=2)
    assert two_rooms.nights == 3 and two_rooms.rooms == 2
    lantern_total = money(stay_total(lantern, two_rooms))
    rows.append(Row(
        "Registered guest", "Booking", results_url(two_rooms),
        [(None, [click_card(lantern["name"])]),
         ("on the details page", [click("Book now", role="link", kind="button", nav=True)]),
         ("on the booking page", [checkpoint(
             "check the booking summary", "Before confirming, on the booking page",
             *summary_shows("Nights: 3", "Rooms: 2", "Total: " + lantern_total))]
          + guest_details() + [confirm])],
        [heading("Booking confirmed"), shows("Total: " + lantern_total)],
        3, 3, stay=two_rooms, final_lead="After clicking Confirm booking", fails_on=["v2"]))

    # Booking form validation
    validation_url = book_url(lantern, sg)
    for name, email, mobile, message, likelihood in (
            ("", GUEST_EMAIL, GUEST_MOBILE, "Enter your full name.", 2),
            (GUEST_NAME, "alex.tester.northwind.test", GUEST_MOBILE, "Enter a valid email address.", 2),
            (GUEST_NAME, GUEST_EMAIL, "71234567", "Enter a valid Singapore mobile number.", 3)):
        rows.append(Row(
            "Registered guest", "Booking validation", validation_url,
            [(None, guest_details(name, email, mobile) + [click("Confirm booking")])],
            [shows(message), heading(BOOK_HEADING)],
            2, likelihood, stay=sg, final_lead="After clicking Confirm booking"))

    # Confirmation and My bookings
    april = Stay("Singapore", "20/04/2027", "22/04/2027")
    kestrel = prop("Kestrel Garden Inn")
    rows.append(Row(
        "Registered guest", "Confirmation", book_url(kestrel, april),
        [(None, guest_details() + [confirm])],
        [heading("Booking confirmed"), shows("Booking reference: NW-100001"), shows(kestrel["name"]),
         shows("Check-in: " + april.checkin), shows("Check-out: " + april.checkout),
         shows("Total: " + money(stay_total(kestrel, april)))],
        3, 2, stay=april, final_lead="After clicking Confirm booking"))

    tidewater = prop("Tidewater Residences")
    rows.append(Row(
        "Registered guest", "My bookings", book_url(tidewater, sg),
        [(None, guest_details() + [confirm]),
         ("on the confirmation page", [click("Go to My bookings", role="link", nav=True)])],
        [heading("My bookings"), shows("NW-100001"), shows(tidewater["name"]),
         shows("%s to %s" % (sg.checkin, sg.checkout)), shows(money(stay_total(tidewater, sg))),
         shows("Confirmed"), button("Cancel booking")],
        3, 2, smoke=True, stay=sg, final_lead="On the My bookings page"))

    rows.append(Row(
        "Registered guest", "My bookings", book_url(lantern, sg),
        [(None, guest_details() + [confirm]),
         ("on the confirmation page", [click("Go to My bookings", role="link", nav=True)]),
         ("on the My bookings page", [click("Cancel booking")])],
        [shows("Cancelled"), not_shows("Confirmed"), button("Cancel booking", present=False)],
        3, 2, stay=sg, final_lead="After clicking Cancel booking"))

    rows.append(Row(
        "Registered guest", "My bookings", "bookings.html?as=tester",
        [], [heading("My bookings"), shows("You have no bookings yet.")],
        1, 1))
    return rows


def feature_rows():
    rows = []
    sg = Stay("Singapore", "12/03/2027", "15/03/2027")
    sg_results = search(sg)
    free = apply_filters(sg_results, free_cancellation=True)
    assert 0 < len(free) < len(sg_results)
    rows.append(Row(
        "Registered guest", "Free cancellation", results_url(sg),
        [(None, [tick("Free cancellation"), click("Apply filters")])],
        [count_is(len(free)), results_are(free)],
        2, 3, stay=sg, final_lead="After clicking Apply filters"))

    tokyo = Stay("Tokyo", "12/03/2027", "15/03/2027")
    tokyo_results = search(tokyo)
    assert any(p["freeCancellation"] for p in tokyo_results) and not all(p["freeCancellation"] for p in tokyo_results)
    rows.append(Row(
        "Registered guest", "Free cancellation", results_url(tokyo),
        [], [card_badge(p["name"], p["freeCancellation"]) for p in tokyo_results],
        2, 2, stay=tokyo, open_note=" and look at each property card"))

    marina = prop("Northwind Marina Suites")
    assert marina["freeCancellation"]
    rows.append(Row(
        "Registered guest", "Free cancellation", results_url(sg),
        [(None, [click_card(marina["name"])])],
        [heading(marina["name"]), shows("Free cancellation")],
        2, 2, stay=sg, final_lead="On the details page"))

    free_hotels = apply_filters(sg_results, types=["Hotel"], free_cancellation=True)
    assert free_hotels
    rows.append(Row(
        "Registered guest", "Free cancellation", results_url(sg),
        [(None, [tick("Free cancellation"), tick("Hotel"), click("Apply filters"),
                 checkpoint("check the results line", "After clicking Apply filters", count_is(len(free_hotels)), results_are(free_hotels))]),
         ("next", [click("Clear filters")])],
        [unticked("Free cancellation"), unticked("Hotel"), count_is(len(sg_results))],
        2, 2, stay=sg, final_lead="After clicking Clear filters"))
    return rows


# --------------------------------------------------------------------------
# Workbooks
# --------------------------------------------------------------------------

HOW_TO_USE = [
    ("title", "Northwind Stays UAT test sheet"),
    ("text", "Northwind Stays is a fictional hotel booking website used as a controlled UAT target. "
             "No real bookings, payments or personal data."),
    ("text", "This sheet tests {build_name} at {build_url}"),
    ("heading", "Columns (worksheet \"Tests\", Excel table \"UATSteps\")"),
    ("col", "StepID", "Unique ID of the test row (NW01, NW02, ...)."),
    ("col", "Persona", "Who is doing the test (for example Registered guest, Locked user, Signed-out visitor)."),
    ("col", "Feature", "The area of the site being tested."),
    ("col", "WhatToDo", "Every action, in order. Starts with the full URL to open. Says exactly which value goes into which field."),
    ("col", "TestData", "The values used in WhatToDo, listed for reference."),
    ("col", "ExpectedResult", "What the screen must show. Text in double quotes must match the page exactly."),
    ("col", "Impact", "How bad a failure would be: 1 (low) to 3 (high)."),
    ("col", "Likelihood", "How likely a failure is: 1 (low) to 3 (high)."),
    ("col", "RiskScore", "Impact multiplied by Likelihood (1 to 9). Stored as a plain number."),
    ("col", "Smoke", "Yes for the quick smoke test rows, otherwise No."),
    ("col", "Status", "Not run, Pass or Fail. Starts as Not run."),
    ("col", "RunDate", "Date the row was run. Empty until run."),
    ("col", "Notes", "What was seen, especially for a Fail. Empty until run."),
    ("heading", "Run order"),
    ("text", "1. Run every row with Smoke = Yes first."),
    ("text", "2. Then run the remaining rows by RiskScore, highest first."),
    ("text", "3. When RiskScore is equal, run the lower StepID first."),
    ("heading", "Each row is self-contained"),
    ("text", "Start every row in a fresh browser session and open the URL at the start of WhatToDo. "
             "No row depends on data from another row, so rows can run in any order."),
    ("text", "Rows that test login or the one-time code use the real login, and the Access control row starts signed out on purpose. "
             "All other rows start on a signed-in page with ?as=tester in the URL. That is a test shortcut that signs in "
             "tester@northwind.test straight away."),
    ("text", "Mark a row Pass only if everything in ExpectedResult is true. Otherwise mark it Fail and describe what you saw in Notes."),
    ("heading", "Test accounts"),
    ("text", "tester@northwind.test / Test@1234: normal user. locked@northwind.test / Test@1234: locked account. One-time code: 123456."),
]


def write_how_to_use(wb, build_name, build_url):
    ws = wb.create_sheet("How to use")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    wrap = Alignment(wrap_text=True, vertical="top")
    row = 1
    for item in HOW_TO_USE:
        kind = item[0]
        if kind == "col":
            ws.cell(row=row, column=1, value=item[1]).font = Font(bold=True)
            ws.cell(row=row, column=2, value=item[2]).alignment = wrap
        else:
            text = item[1].format(build_name=build_name, build_url=build_url)
            cell = ws.cell(row=row, column=1, value=text)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            cell.alignment = wrap
            if kind == "title":
                cell.font = Font(bold=True, size=14)
            elif kind == "heading":
                cell.font = Font(bold=True, size=12)
                ws.row_dimensions[row].height = 24
            elif len(text) > 110:
                ws.row_dimensions[row].height = 15 * (len(text) // 110 + 1)
        row += 1


def write_workbook(path, rows, base, build, build_name):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tests"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([
            r.step_id, r.persona, r.feature, r.what_to_do(base, build), r.test_data(), r.expected_result(),
            r.impact, r.likelihood, r.impact * r.likelihood, "Yes" if r.smoke else "No", "Not run", None, None,
        ])

    last_col = get_column_letter(len(COLUMNS))
    table = Table(displayName="UATSteps", ref="A1:%s%d" % (last_col, len(rows) + 1))
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=1, max_row=len(rows) + 1, max_col=len(COLUMNS)):
        for cell in row:
            cell.alignment = wrap
    for i, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[name]
    ws.freeze_panes = "A2"

    col = {name: get_column_letter(i) for i, name in enumerate(COLUMNS, start=1)}
    span = lambda name: "%s2:%s%d" % (col[name], col[name], VALIDATION_LAST_ROW)
    for name in ("Impact", "Likelihood"):
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid " + name, error=name + " must be a whole number from 1 to 3.")
        dv.add(span(name))
        ws.add_data_validation(dv)
    for name, choices in (("Smoke", "Yes,No"), ("Status", "Pass,Fail,Not run")):
        dv = DataValidation(type="list", formula1='"%s"' % choices, allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid " + name, error=name + " must be one of: " + choices.replace(",", ", ") + ".")
        dv.add(span(name))
        ws.add_data_validation(dv)

    write_how_to_use(wb, build_name, "%s/%s/" % (base, build))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--base-url", required=True, help="for example https://USERNAME.github.io/northwind-stays")
    parser.add_argument("--out", default=str(ROOT / "sheets"), help="folder for the .xlsx files")
    parser.add_argument("--json", default=str(ROOT / "verify" / "test_rows.json"), help="where to write the rows for verify/verify.js")
    args = parser.parse_args()

    check_data_rules()
    base = args.base_url.rstrip("/")
    core = core_rows()
    extra = feature_rows()
    for number, row in enumerate(core + extra, start=1):
        row.step_id = "NW%02d" % number

    smoke = [r for r in core if r.smoke]
    assert len(smoke) == 6, "exactly 6 smoke rows required, found %d" % len(smoke)
    assert not any(r.fails_on for r in smoke), "smoke rows must not be affected by Build 1.1.0 changes"
    assert sum(1 for r in core if r.fails_on) == 3

    out = Path(args.out)
    write_workbook(out / "Northwind_Tests_Build_1.0.0.xlsx", core, base, "v1", "Build 1.0.0")
    write_workbook(out / "Northwind_Tests_Build_1.1.0.xlsx", core, base, "v2", "Build 1.1.0")
    write_workbook(out / "Northwind_NewFeature_Rows_1.1.0.xlsx", extra, base, "v2", "Build 1.1.0 (new feature rows)")

    payload = {
        "baseUrl": base,
        "core": [r.to_json(base, "v1") for r in core],
        "feature": [r.to_json(base, "v2") for r in extra],
    }
    json_path = Path(args.json)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")

    print("Core rows: %d (%s to %s), smoke rows: %s" % (len(core), core[0].step_id, core[-1].step_id, ", ".join(r.step_id for r in smoke)))
    print("New feature rows: %d (%s to %s)" % (len(extra), extra[0].step_id, extra[-1].step_id))
    for name in ("Northwind_Tests_Build_1.0.0.xlsx", "Northwind_Tests_Build_1.1.0.xlsx", "Northwind_NewFeature_Rows_1.1.0.xlsx"):
        print("Wrote", out / name)
    print("Wrote", json_path)


if __name__ == "__main__":
    main()
