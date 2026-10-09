#!/usr/bin/env python3
"""Check the structure of the three Northwind Stays test workbooks.

    python verify/check_sheets.py --base-url https://USERNAME.github.io/northwind-stays

Prints a JSON summary ({"ok": bool, "checks": [...]}) for verify/verify.js.
"""

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
COLUMNS = ["StepID", "Persona", "Feature", "WhatToDo", "TestData", "ExpectedResult",
           "Impact", "Likelihood", "RiskScore", "Smoke", "Status", "RunDate", "Notes"]
BOOKS = [
    ("Northwind_Tests_Build_1.0.0.xlsx", "v1", True),
    ("Northwind_Tests_Build_1.1.0.xlsx", "v2", True),
    ("Northwind_NewFeature_Rows_1.1.0.xlsx", "v2", False),
]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", required=True)
    parser.add_argument("--sheets", default=str(ROOT / "sheets"))
    args = parser.parse_args()
    base = args.base_url.rstrip("/")
    checks = []

    def check(name, ok, detail=""):
        checks.append({"name": name, "ok": bool(ok), "detail": detail})

    rows_by_book = {}
    for filename, build, is_core in BOOKS:
        path = Path(args.sheets) / filename
        if not path.exists():
            check(filename + " exists", False, "missing")
            continue
        wb = load_workbook(path)
        check(filename + ": worksheets", wb.sheetnames == ["Tests", "How to use"], ", ".join(wb.sheetnames))
        ws = wb["Tests"]
        header = [c.value for c in ws[1]]
        check(filename + ": columns in order", header == COLUMNS, ", ".join(str(h) for h in header))
        tables = ws.tables
        table = tables.get("UATSteps")
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2, values_only=False)]
        rows = [r for r in rows if any(v is not None for v in r)]
        rows_by_book[filename] = rows
        expected_ref = "A1:M%d" % (len(rows) + 1)
        check(filename + ": table UATSteps", table is not None and table.ref == expected_ref,
              (table.ref if table else "no table") + ", expected " + expected_ref)
        check(filename + ": frozen header row", ws.freeze_panes == "A2", str(ws.freeze_panes))
        check(filename + ": wrapped text", all(ws.cell(row=r, column=c).alignment.wrap_text
                                                for r in range(1, len(rows) + 2) for c in range(1, 14)))

        ids = [r[0] for r in rows]
        numbers = [int(i[2:]) for i in ids if re.fullmatch(r"NW\d{2}", str(i))]
        check(filename + ": StepIDs NWnn in sequence", len(numbers) == len(ids) and numbers == list(range(numbers[0], numbers[0] + len(ids))),
              "%s to %s" % (ids[0], ids[-1]))
        bad_risk = [r[0] for r in rows if not (isinstance(r[6], int) and isinstance(r[7], int) and 1 <= r[6] <= 3 and 1 <= r[7] <= 3
                                                and isinstance(r[8], int) and r[8] == r[6] * r[7])]
        check(filename + ": Impact/Likelihood 1-3, RiskScore = plain number Impact x Likelihood", not bad_risk, ", ".join(bad_risk))
        formulas = [r[0] for r in rows if any(isinstance(v, str) and v.startswith("=") for v in r)]
        check(filename + ": no formulas", not formulas, ", ".join(formulas))
        check(filename + ": Smoke is Yes/No", all(r[9] in ("Yes", "No") for r in rows))
        check(filename + ": Status starts as Not run, RunDate and Notes empty",
              all(r[10] == "Not run" and r[11] is None and r[12] is None for r in rows))
        smoke = [r[0] for r in rows if r[9] == "Yes"]
        if is_core:
            check(filename + ": exactly 6 smoke rows", len(smoke) == 6, ", ".join(smoke))
        prefix = "Open %s/%s/" % (base, build)
        wrong_url = [r[0] for r in rows if not str(r[3]).startswith(prefix)]
        check(filename + ": every row starts with " + prefix, not wrong_url, ", ".join(wrong_url))
        other = "/v2/" if build == "v1" else "/v1/"
        check(filename + ": no row points at the other build", not any(other in str(r[3]) for r in rows))
        years = {}
        for r in rows:
            for y in re.findall(r"\b\d{1,2}/\d{1,2}/(\d{4})\b", str(r[3]) + " " + str(r[4]) + " " + str(r[5])):
                years.setdefault(y, set()).add(r[0])
        non_2027 = {y: sorted(v) for y, v in years.items() if y != "2027"}
        past_rows = sorted(set().union(*non_2027.values())) if non_2027 else []
        check(filename + ": all dates in 2027 except the single past-date row", len(past_rows) <= 1,
              "non-2027 dates in: " + (", ".join(past_rows) or "none"))
        shortcut_missing = [r[0] for r in rows if not str(r[3]).startswith(prefix + "login.html")
                            and r[2] != "Access control" and "?as=tester" not in str(r[3])]
        check(filename + ": non-login rows start on a signed-in page with ?as=tester", not shortcut_missing, ", ".join(shortcut_missing))

        validations = {}
        for dv in ws.data_validations.dataValidation:
            for rng in str(dv.sqref).split():
                validations[rng[0]] = (dv.type, dv.formula1, dv.formula2)
        check(filename + ": Impact validation whole 1-3", validations.get("G") == ("whole", "1", "3"), str(validations.get("G")))
        check(filename + ": Likelihood validation whole 1-3", validations.get("H") == ("whole", "1", "3"), str(validations.get("H")))
        check(filename + ": Smoke validation Yes/No list", validations.get("J") == ("list", '"Yes,No"', None), str(validations.get("J")))
        check(filename + ": Status validation Pass/Fail/Not run list", validations.get("K") == ("list", '"Pass,Fail,Not run"', None), str(validations.get("K")))
        how = wb["How to use"]
        text = " ".join(str(c.value) for row in how.iter_rows() for c in row if c.value)
        check(filename + ": How to use explains run order and self-contained rows",
              "Smoke = Yes" in text and "RiskScore, highest first" in text and "lower StepID first" in text and "self-contained" in text)

    v1 = rows_by_book.get(BOOKS[0][0], [])
    v2 = rows_by_book.get(BOOKS[1][0], [])
    same = len(v1) == len(v2) and all(
        [a for i, a in enumerate(r1) if i != 3] == [b for i, b in enumerate(r2) if i != 3]
        and r1[3].replace("/v1/", "/v2/") == r2[3] for r1, r2 in zip(v1, v2))
    check("Build 1.0.0 and 1.1.0 sheets have identical rows apart from the build in the URL", same)
    feature = rows_by_book.get(BOOKS[2][0], [])
    if v1 and feature:
        check("New feature rows continue the StepID numbering",
              int(feature[0][0][2:]) == int(v1[-1][0][2:]) + 1, "%s follows %s" % (feature[0][0], v1[-1][0]))
        check("3 or 4 new feature rows", 3 <= len(feature) <= 4, str(len(feature)))
    if v1:
        check("About 30 core rows", 28 <= len(v1) <= 36, str(len(v1)))

    json.dump({"ok": all(c["ok"] for c in checks), "checks": checks}, sys.stdout, indent=1)


if __name__ == "__main__":
    main()
